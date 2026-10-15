"""把课程仓库里的真实资料整理成知识库 knowledge/*.md。

资料来源（均在本仓库 2026-capstone/ 下）：
- 各周课件 .pptx：逐页抽取文字和表格，每页写成一个 "## 第 N 页 标题" 小节；
- 各周 README.md / assignment.md、DGX Spark 登录手册；
- 我自己前三周的 提交/提交说明.md（实验记录，便于问"我当时的结果"）。

生成的 knowledge/ 已随作业提交，复现 Agent 时无需再运行本脚本（需要 python-pptx）。
用法：python build_kb.py [--course ../..]
"""
import argparse
import re
from pathlib import Path

HERE = Path(__file__).resolve().parent

# (输出文件名, 相对 2026-capstone 的路径, 标题)
SOURCES = [
    ("w01-slides", "week01-llm-inference/llm-inference-tutorial.pptx", "第一周课件：LLM 推理"),
    ("w01-readme", "week01-llm-inference/README.md", "第一周 README：LLM 推理"),
    ("w01-assignment", "week01-llm-inference/assignment.md", "第一周作业要求"),
    ("w02-slides-app", "week02-llm-application/llm-application-basics.pptx", "第二周课件：LLM 应用开发基础"),
    ("w02-slides-structured", "week02-llm-application/structured-output-and-result-processing.pptx",
     "第二周课件：结构化输出与结果处理"),
    ("w02-readme", "week02-llm-application/README.md", "第二周 README：LLM 应用开发"),
    ("w02-assignment", "week02-llm-application/assignment.md", "第二周作业要求"),
    ("w03-slides", "week03-agent-tool-calling/minimal-agent-and-tool-calling.pptx",
     "第三周课件：Minimal Agent 与 Tool Calling"),
    ("w03-readme", "week03-agent-tool-calling/README.md", "第三周 README：Agent 与 Tool Calling"),
    ("w03-assignment", "week03-agent-tool-calling/assignment.md", "第三周作业要求"),
    ("w05-slides", "week05-agent-memory-tools/agent-memory-tools-and-mcp.pptx",
     "第五周课件：Agent Memory、Tool 与 MCP"),
    ("w05-readme", "week05-agent-memory-tools/README.md", "第五周 README：Agent Memory、Tool 与 MCP"),
    ("w06-slides-rag", "week06-1-knowledge-rag/knowledge-and-rag.pptx", "第六周第 1 次课课件：Knowledge 与 RAG"),
    ("w06-readme-rag", "week06-1-knowledge-rag/README.md", "第六周第 1 次课 README：Knowledge 与 RAG"),
    ("w06-assignment", "week06-1-knowledge-rag/assignment.md", "第六周作业要求：完成 Agent v1"),
    ("w06-slides-serving", "week06-2-dgx-local-serving/dgx-local-serving.pptx",
     "第六周第 2 次课课件：DGX Spark 本地部署与 vLLM"),
    ("w06-readme-serving", "week06-2-dgx-local-serving/README.md", "第六周第 2 次课 README：DGX Spark 与 vLLM"),
    ("dgx-login", "dgx-spark-login.md", "DGX Spark 登录手册"),
    ("my-w01-report", "week01-llm-inference/提交/提交说明.md", "我的第一周作业提交说明（实验记录）"),
    ("my-w02-report", "week02-llm-application/提交/提交说明.md", "我的第二周作业提交说明（实验记录）"),
    ("my-w03-report", "week03-agent-tool-calling/提交/提交说明.md", "我的第三周作业提交说明（实验记录）"),
]


def shape_lines(shape):
    """递归取出形状中的文字；表格按行用 | 连接。"""
    if shape.shape_type == 6:  # GROUP
        for sub in shape.shapes:
            yield from shape_lines(sub)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                yield " | ".join(cells)
    elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text:
            yield re.sub(r"[ \t]*\n[ \t]*", " ", text)


def pptx_to_md(path, title):
    from pptx import Presentation
    out = [f"# {title}", "", f"原始文件：`{path.name}`", ""]
    for number, slide in enumerate(Presentation(path).slides, 1):
        lines = [line for shape in slide.shapes for line in shape_lines(shape)]
        # 去掉单独的页码和重复行，保留顺序
        seen, body = set(), []
        for line in lines:
            if line.isdigit() or line in seen:
                continue
            seen.add(line)
            body.append(line)
        if not body:
            continue
        heading = body[0][:60]
        out += [f"## 第 {number} 页 {heading}", ""] + body[1:] + [""]
    return "\n".join(out)


def md_to_md(path, title):
    text = path.read_text(encoding="utf-8")
    # 去掉图片和相对链接的目标，只保留可读文字
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    text = re.sub(r"\[([^\]]+)\]\((?!https?://)[^)]*\)", r"\1", text)
    lines = text.splitlines()
    if lines and lines[0].startswith("# "):
        lines = lines[1:]
    return "\n".join([f"# {title}", "", f"原始文件：`{path.as_posix()}`", ""] + lines).strip() + "\n"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--course", type=Path, default=HERE.parents[1], help="2026-capstone 目录")
    parser.add_argument("--out", type=Path, default=HERE / "knowledge")
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    for name, rel, title in SOURCES:
        path = args.course / rel
        if path.suffix == ".pptx":
            md = pptx_to_md(path, title)
        else:
            md = md_to_md(path, title).replace(f"原始文件：`{path.as_posix()}`",
                                               f"原始文件：`2026-capstone/{rel}`")
        (args.out / f"{name}.md").write_text(md, encoding="utf-8")
        print(f"{name}.md  {len(md):>6} 字符  ← {rel}")


if __name__ == "__main__":
    main()
