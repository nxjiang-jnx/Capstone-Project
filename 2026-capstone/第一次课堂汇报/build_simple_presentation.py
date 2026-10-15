from pathlib import Path
import io, json, hashlib, math, zipfile
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
import pymupdf as fitz

OUT=Path(__file__).resolve().parent
C=OUT.parent
W1=C/'第一周_LLM推理/提交'
W2=C/'第二周_LLM应用开发/提交'
W3=C/'第三周_Agent与工具调用/提交'
FONT='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
BOLD='/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc'
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='前三周课程学习汇报'
prs.core_properties.author='姜楠翔'
slides=[];checks=[]

class Slide:
    def __init__(self,title,seconds=0,source='',appendix=False,cover=False):
        self.title=title;self.seconds=seconds;self.source=source;self.appendix=appendix
        self.slide=prs.slides.add_slide(prs.slide_layouts[6])
        self.slide.background.fill.solid();self.slide.background.fill.fore_color.rgb=RGBColor(255,255,255)
        self.img=Image.new('RGB',(1920,1080),'white');self.draw=ImageDraw.Draw(self.img)
        slides.append(self)
        if not cover:
            self.text(.9,.68,11.5,1.0,title,32,True,color='123A56')
            self.rect(.92,1.48,.7,.045,'006B75')
    def rect(self,x,y,w,h,color):
        shape=self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(color);shape.line.fill.background()
        self.draw.rectangle(tuple(round(v*144) for v in (x,y,x+w,y+h)),fill='#'+color)
    def node(self,x,y,w,h,label):
        shape=self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        shape.adjustments[0]=.13;shape.fill.solid();shape.fill.fore_color.rgb=RGBColor(255,255,255)
        shape.line.color.rgb=RGBColor.from_string('007B88');shape.line.width=Pt(1.3)
        self.draw.rounded_rectangle(tuple(round(v*144) for v in (x,y,x+w,y+h)),radius=12,fill='white',outline='#007B88',width=3)
        self.text(x+.64,y+.17,w-1.1,h-.1,label,20,color='006B75')
    def arrow(self,x1,y1,x2,y2,color='007B88'):
        shape=self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
        shape.line.color.rgb=RGBColor.from_string(color);shape.line.width=Pt(1.3)
        ln=shape._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln')
        end=OxmlElement('a:tailEnd');end.set('type','triangle');end.set('w','sm');end.set('len','sm');ln.append(end)
        pts=(round(x1*144),round(y1*144),round(x2*144),round(y2*144));self.draw.line(pts,fill='#'+color,width=3)
        angle=math.atan2(y2-y1,x2-x1);tip=(x2*144,y2*144)
        wing=[tip,(tip[0]-12*math.cos(angle)+5*math.sin(angle),tip[1]-12*math.sin(angle)-5*math.cos(angle)),(tip[0]-12*math.cos(angle)-5*math.sin(angle),tip[1]-12*math.sin(angle)+5*math.cos(angle))]
        self.draw.polygon(wing,fill='#'+color)
    def text(self,x,y,w,h,content,size=24,bold=False,mono=False,color='151515'):
        f=ImageFont.truetype(BOLD if bold else FONT,size*2,index=7 if mono else 2)
        lines=[]
        for para in content.split('\n'):
            curr=''
            for char in para:
                if self.draw.textlength(curr+char,font=f)>w*144-6 and curr:
                    lines.append(curr);curr=char
                else:curr+=char
            lines.append(curr)
        step=size/72*1.4
        if len(lines)*step>h+.05:checks.append({'slide':len(slides),'overflow':content[:60]})
        for i,line in enumerate(lines):
            yy=y+i*step
            shape=self.slide.shapes.add_textbox(Inches(x),Inches(yy),Inches(w),Inches(step+.03))
            tf=shape.text_frame;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0;tf.word_wrap=False
            p=tf.paragraphs[0];p.text=line;p.font.name='Noto Sans Mono CJK SC' if mono else 'Noto Sans CJK SC';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
            self.draw.text((round(x*144),round(yy*144)),line,font=f,fill='#'+color,anchor='lt')
    def items(self,items):
        for i,line in enumerate(items):
            y=2.05+i*1.25
            self.rect(.95,y+.15,.055,.055,'006B75')
            self.text(1.2,y,11.05,1.0,line)
            if i<len(items)-1:self.rect(1.2,y+.99,10.95,.007,'B3DBE5')
    def table(self,headers,rows,widths,top=2.0,size=22):
        xs=[.98]
        for w in widths[:-1]:xs.append(xs[-1]+w)
        self.rect(.9,top-.17,sum(widths)+.1,.78,'FFFFFF')
        for j,t in enumerate(headers):self.text(xs[j],top,widths[j]-.16,.8,t,size,True)
        for i,row in enumerate(rows):
            for j,t in enumerate(row):self.text(xs[j],top+.86+i*.72,widths[j]-.16,.8,t,size)
            self.rect(.9,top+1.42+i*.72,sum(widths)+.1,.008,'B3DBE5')
    def notes(self,speech):
        self.speech=speech
        meta='答疑附录，不计入正文时长' if self.appendix else f'建议 {self.seconds} 秒'
        self.slide.notes_slide.notes_text_frame.text=f'{self.title}\n{meta}\n\n【讲稿】\n{speech}\n\n【资料依据，不用念】\n{self.source}'

s=Slide('前三周学习与作业汇报',15,cover=True)
s.text(.95,.85,8.0,.7,'前三周学习与作业汇报',23,color='151515')
s.text(.92,1.88,8.0,2.0,'代码生成的\n自动测试与反馈修复',42,True,color='123A56')
s.text(.96,4.02,7.8,.8,'模型推理 · 代码助手 · Agent',23,color='151515')
s.rect(.97,5.3,.8,.045,'006B75')
s.text(.96,5.85,7.8,.8,'姜楠翔    23230610',22,color='123A56')
s.node(9.37,2.13,2.4,.78,'生成代码')
s.node(9.37,3.49,2.4,.78,'执行测试')
s.node(9.37,4.85,2.4,.78,'反馈修复')
s.arrow(10.57,2.96,10.57,3.39)
s.arrow(10.57,4.32,10.57,4.75)
s.rect(12.22,2.52,.012,2.73,'007B88')
s.rect(11.82,5.23,.40,.012,'007B88')
s.arrow(12.22,2.52,11.87,2.52)
s.notes('大家好，我是姜楠翔。前三周我做了模型推理、代码助手和 Agent 实验。这次重点汇报一个小拓展：给生成的代码加上测试，再让模型根据错误自动修改。')

d=json.loads((W1/'run.json').read_text())
s=Slide('第一周：Qwen 模型推理',40,'第一周/提交/run.json；qwen_inference.py:58–95、139–147。最后层取样为[0,-1,:4]。')
s.items(['Qwen2.5-1.5B：单次推理 3.622 秒，生成 74 个 token', '最后层 hidden state：[1, 48, 1536] → [1, 1, 1536]', '部分数值：0.31055，2.0，0.0，1.58594'])
s.notes('第一周，我用 Transformers 在 GPU 上跑 Qwen 一点五 B。先预热一次，再正式计时，并等待 GPU 完成计算，避免计时偏差。一次推理用了三点六二二秒，生成七十四个新 token，不把输入算进去。我还观察了最后一层的数据：开始处理四十八个输入位置，之后带着缓存，每次只处理一个新位置。不过，模型对缓存的解释不准确，说明跑通和答对还要分开检查。')

s=Slide('第二周：我想解决什么问题',30,'第二周/提交/results/plain.json、example.json、example_chat.turn01/02.json；原示例未执行功能测试。')
s.items(['已跑通普通回答、流式输出和结构化校验', '原示例检查了格式和语法，没有执行代码测试', '我的拓展：自动测试 → 反馈具体错误 → 自动重试'])
s.notes('第二周，我跑通了普通回答、流式输出和结构化校验。流式就是边生成边显示，结构化则要求按统一的 JSON 格式回答，方便程序读取。但格式和语法正确，不代表代码结果正确。所以我加上测试用例，让程序实际运行代码，失败后自动反馈，不用每次手动追问。')

s=Slide('代码改动：校验结果决定是否重试',40,'llm_app.py:123–168、325–338。代码省略异常处理和函数名检查；JSON、Schema、语法、函数名、测试共五层校验。')
s.text(1.05,1.75,11.25,.7,'JSON → 字段 → 语法 → 函数名 → 测试',24)
s.rect(.94,2.48,11.45,2.03,'FFFFFF')
s.rect(.94,2.48,.025,2.03,'007B88')
s.text(1.05,2.65,11.25,1.95,'''result = validate_text(text, task)
feedback = repair_feedback(result, task)
if feedback and repairs < args.repair:
    repairs, source = repairs + 1, "auto_repair"''',22,mono=True)
s.text(1.05,4.85,11.2,.7,'history：保留错误回答，再追加反馈',24)
s.text(1.05,5.8,11.2,.7,'fresh：丢弃错误回答，用原任务和反馈重新生成',24)
s.notes('这里主要加了两部分。第一部分是校验：检查 JSON、字段、语法和函数名，再执行测试，比较返回值和类型。测试放在独立进程里，超过五秒就停止，避免卡住。第二部分是修复：把具体错误发给模型，最多重试三次。我比较两种方式：history 保留错误回答继续改；fresh 丢弃错误回答，只给原任务和错误提示，让模型重新写。')

s=Slide('实验：保留错误回答，还是重新生成？',35,'results/ext_v2_*、ext_v3fresh_*；次数为自动修复次数，不含首次生成；上限均为3次。')
s.table(['任务','保留错误回答','丢弃错误回答后重试'],[
    ['正数计数','3 次修复，通过','1 次修复，通过'],
    ['第二大不同值','3 次修复，失败','2 次修复，通过'],
    ['游程编码','3 次修复，失败','3 次修复，失败']
], [3.0,3.8,4.55],size=22)
s.text(1.05,5.75,11.2,.8,'通过的两项任务分别完成 5/5、7/7 个测试用例。',24)
s.notes('实验用了三个任务，模型和生成设置保持一致，只比较两种修复方式。保留错误回答时，正数计数修复三次才通过，另外两项失败。重新生成后，正数计数一次修复通过，第二大不同值两次修复通过，分别通过五个和七个测试。游程编码仍然失败。因此在这组实验里，重新生成修好了更多任务，也减少了重试，但还需要更多任务验证。')

runpaths=sorted((W3/'results').glob('run*.json'))
runs=[json.loads(p.read_text()) for p in runpaths]
s=Slide('第三周：Agent 自测与五组实验',35,'agent_v0.py --selftest：本次复核14项全PASS；run1–run5。组3→4同时改变模型与采样开关。')
s.text(1.05,1.68,11.2,.8,'自测 14 项全部 PASS；固定任务的正确答案为 94%。',24)
s.table(['组','模型和解码','status','最终答对'],[
    ['1','1.5B，贪心',runs[0]['status'],'否'],
    ['2','1.5B，采样 T=0.7',runs[1]['status'],'否'],
    ['3','1.5B，采样 T=1.2',runs[2]['status'],'否'],
    ['4','3B，贪心',runs[3]['status'],'否'],
    ['5','7B，贪心',runs[4]['status'],'否']
], [.75,4.55,4.5,1.55],top=2.6,size=21)
s.notes('第三周，我让 Agent 读取测试报告，再用工具计算通过率。报告有四十七项通过、三项失败，答案应是百分之九十四。程序的十四项自测全通过，但换模型大小和生成方式后，五组实验最终都没答对。两组虽然显示 ok，也只是正常结束。另外，第三组到第四组同时改了模型和采样方式，这部分还需要补成严格的单变量对照。')

s=Slide('看调用链：问题出在哪一步',45,'run1_1.5B_greedy.json、run4_3B_greedy.json；ext_3B/7B_greedy_v2/v3。每条链的逐步详情在讲稿中。')
s.items(['1.5B：错误路径 → 读取失败 → final 说信息不足', '3B：读文件 → 算出 94 → 回答没有 final JSON', '补充文件名和调用示例后，3B、7B 的 4 次运行均答对'])
s.notes('我查看了两条调用链，也就是模型每一步说了什么、调用了什么、工具返回了什么。一点五 B 把路径写错，工具没读到文件，最后却正常结束。三 B 读到了文件，计算器也算出九十四，但它没按要求输出 final JSON，程序不接受，最后用完了步数。我据此补充文件名、收尾规则和调用示例。三 B、七 B 在两版新提示下的四次运行都答对了，但一点五 B 仍会写错计算器参数。')

s=Slide('这次拓展的结果与不足',20,'结论仅限现有三个代码任务；单变量补跑与最终答案检查为下一步，尚未实施。')
s.items(['代码助手增加了测试和自动反馈，减少人工追问', '重新生成在本组任务中效果更好，游程编码仍失败', '下一步补严单变量实验，并检查 Agent 的最终答案'])
s.notes('这次拓展把代码检查和错误反馈连了起来，重新生成在这几个任务中效果更好。Agent 实验也说明，调用工具不代表一定做对。接下来会补严对照实验，并检查最终答案，而不是只看是否正常结束。谢谢大家。')

assert sum(s.seconds for s in slides)==260
assert len(slides)==8
assert not checks,checks
OUT.mkdir(exist_ok=True);(OUT/'预览').mkdir(exist_ok=True)
for p in (OUT/'预览').glob('*.png'):p.unlink()
name='23230610_姜楠翔_前三周学习汇报'
prs.save(OUT/(name+'.pptx'))
pdf=fitz.open()
for i,s in enumerate(slides,1):
    s.img.save(OUT/'预览'/f'{i:02d}.png')
    buff=io.BytesIO();s.img.save(buff,format='JPEG',quality=92)
    page=pdf.new_page(width=960,height=540);page.insert_image(page.rect,stream=buff.getvalue())
pdf.save(OUT/(name+'_预览.pdf'),deflate=True)
contact=Image.new('RGB',(1920,math.ceil(len(slides)/3)*390),'#FFFFFF');draw=ImageDraw.Draw(contact)
small=ImageFont.truetype(FONT,18,index=2)
for i,s in enumerate(slides):
    x=10+i%3*640;y=12+i//3*390;contact.paste(s.img.resize((620,349)),(x,y));draw.text((x,y+355),f'{i+1}. {s.title}',font=small,fill='black')
contact.save(OUT/'预览总览.jpg',quality=92)
md=['# 逐页中文讲稿','','姓名：姜楠翔；学号：23230610。','','共 8 页，建议约 4 分 20 秒。无附录，讲稿已写入 PPT 备注。','','核心问题：代码生成出错后，如何自动检查并修复？保留错误回答，还是用原任务和错误反馈重新生成更有效？','','讲稿补充必要的技术细节，每页只讲一段。表格和代码不必逐行念，正常语速配合指图，控制在课程要求的 3–5 分钟内。']
elapsed=0
for i,s in enumerate(slides,1):
    md.extend(['',f'## 第 {i} 页：{s.title}',''])
    if s.appendix:md.extend(['答疑时使用，不计入正文时间。',''])
    else:
        md.extend([f'建议 {s.seconds} 秒，累计 {(elapsed+s.seconds)//60} 分 {(elapsed+s.seconds)%60} 秒。','']);elapsed+=s.seconds
    md.append(s.speech)
(OUT/'逐页中文讲稿.md').write_text('\n'.join(md)+'\n')
doc=fitz.open();note_pages=[];note_image=None;y=0
for line in md:
    fs=16 if line.startswith('# ') else 13 if line.startswith('## ') else 11
    content=line.lstrip('# ') if line.startswith('#') else line
    font=ImageFont.truetype(FONT,fs*3,index=2);wrapped=[];cur=''
    for ch in content:
        if font.getlength(cur+ch)/3>505 and cur:wrapped.append(cur);cur=ch
        else:cur+=ch
    wrapped.append(cur)
    for part in wrapped:
        if note_image is None or y>780:
            note_image=Image.new('RGB',(1785,2526),'white');note_draw=ImageDraw.Draw(note_image);note_pages.append(note_image);y=45
        note_draw.text((135,round(y*3)),part,font=font,fill='black',anchor='lt');y+=fs*1.7
for im in note_pages:
    buf=io.BytesIO();im.save(buf,format='JPEG',quality=93)
    page=doc.new_page(width=595,height=842);page.insert_image(page.rect,stream=buf.getvalue())
doc.save(OUT/'逐页中文讲稿.pdf',deflate=True)
sources=[]
for root in [W1,W2,W3]:
    for p in sorted(root.rglob('*')):
        if p.is_file() and p.suffix in ['.py','.json','.md','.txt']:
            sources.append({'path':str(p.relative_to(C)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
manifest={'main_slides':8,'appendix_slides':0,'planned_seconds':260,'style':'white background; black and navy text with teal rules; no gray text or fills; large cover title and editable feedback-loop diagram; consistent body typography and table rules; no appendix or footers','sources':sources,'layout_warnings':checks,'slides':[{'number':i,'title':s.title,'source':s.source} for i,s in enumerate(slides,1)],'preview_method':'Pillow renders shared layout; not an Office engine export'}
(OUT/'证据与制作校验.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(json.dumps({'slides':len(slides),'seconds':260,'layout_warnings':checks},ensure_ascii=False))
